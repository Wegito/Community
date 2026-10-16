from datetime import UTC, date, datetime, timedelta

import pytest
from openpyxl import Workbook, load_workbook
from sqlalchemy import select

from community.core.org.models import MemberGroup
from community.core.org.services import add_to_group
from community.core.platform.qr import tokens_for
from community.core.time import to_local, utcnow
from community.extensions import db
from community.modules.laundry.models import LaundryMachine, LaundryPrice, LaundryUse
from community.modules.laundry.services import (
    LaundryError,
    import_google_form_xlsx,
    month_bounds,
    monthly_summary,
    record_use,
)

from .factories import make_user


def make_resident(email='bewohner@example.org', name='Ben Bewohner', **kwargs):
    """A Hausbewohner: may enter laundry. rank=None makes a resident without membership."""
    user = make_user(email=email, name=name, **kwargs)
    group = db.session.scalar(select(MemberGroup).where(MemberGroup.name == 'Hausbewohner'))
    add_to_group(group, user.member, date(2020, 1, 1))
    db.session.commit()
    return user


def form(**overrides):
    data = {'washes': '1', 'dry_program': '', 'container': '', 'submission_id': 'sub-1'}
    data.update(overrides)
    return data


def entries():
    return db.session.scalars(select(LaundryUse).order_by(LaundryUse.id)).all()


def test_member_pays_member_price_for_washing_and_drying(client, login):
    user = make_resident()
    login(user)
    page = client.get('/waesche/neu').get_data(as_text=True)
    assert '0,25 €' in page and 'Schranktrocken +' in page
    response = client.post('/waesche/neu', data=form(washes='2', dried='y', dry_program='Schranktrocken +'))
    assert response.status_code == 302
    (use,) = entries()
    assert (use.washes, use.dried, use.tier) == (2, True, 'member')
    assert use.total_cents == 2 * 25 + 25
    assert '0,75 €' in client.get('/waesche/').get_data(as_text=True)


def test_resident_without_membership_pays_resident_price(client, login):
    user = make_resident(rank=None)
    login(user)
    client.post('/waesche/neu', data=form(washes='1', dried='y', dry_program='Extra Trocken'))
    (use,) = entries()
    assert use.tier == 'resident'
    assert use.total_cents == 150 + 150


def test_double_submit_saves_once(client, login):
    login(make_resident())
    client.post('/waesche/neu', data=form())
    response = client.post('/waesche/neu', data=form(), follow_redirects=True)
    assert len(entries()) == 1
    assert 'schon gespeichert' in response.get_data(as_text=True)


@pytest.mark.parametrize(
    ('data', 'message'),
    [
        (form(washes='0'), 'Waschen oder Trocknen'),
        (form(washes='0', dried='y'), 'Trocknerprogramm'),
        (form(washes='0', dried='y', dry_program='Fön'), 'keine gültige'),
    ],
)
def test_form_validation(client, login, data, message):
    login(make_resident())
    response = client.post('/waesche/neu', data=data)
    assert response.status_code == 200
    assert message.lower() in response.get_data(as_text=True).lower()
    assert entries() == []


def test_only_hausbewohner_may_enter_laundry(client, login):
    login(make_user())  # a Bursch who does not live in the house
    assert client.get('/waesche/neu').status_code == 403


def test_undo_within_minutes_only_for_own_entry(app, client, login):
    user = make_resident()
    other = make_resident(email='other@example.org', name='Other Bewohner')
    login(user)
    client.post('/waesche/neu', data=form())
    (use,) = entries()
    client.post('/logout')
    login(other)
    assert client.post(f'/waesche/{use.id}/rueckgaengig').status_code == 403
    client.post('/logout')
    login(user)
    assert client.post(f'/waesche/{use.id}/rueckgaengig').status_code == 302
    db.session.refresh(use)
    assert use.is_void


def test_undo_window_closes(client, login):
    user = make_resident()
    login(user)
    client.post('/waesche/neu', data=form())
    (use,) = entries()
    use.created_at = utcnow() - timedelta(minutes=11)
    db.session.commit()
    assert client.post(f'/waesche/{use.id}/rueckgaengig').status_code == 403


def test_old_entries_keep_their_price(app):
    user = make_resident()
    first, _ = record_use(user.member, washes=1, dried=False, submission_id='a')
    db.session.add(LaundryPrice(tier='member', kind='wash', price_cents=40, valid_from=date(2026, 1, 2)))
    db.session.commit()
    second, _ = record_use(user.member, washes=1, dried=False, submission_id='b')
    assert first.used_at < second.used_at or first.id < second.id
    assert second.wash_price_cents == 40


def test_missing_price_is_an_error(app):
    user = make_resident()
    with pytest.raises(LaundryError, match='kein Preis'):
        record_use(user.member, washes=1, dried=False, submission_id='x', used_at=datetime(2025, 6, 1, tzinfo=UTC))


def test_monthly_overview_and_excel(client, login):
    user = make_resident()
    resident = make_resident(email='r@example.org', name='Rita Resident', rank=None)
    now = utcnow()
    record_use(user.member, washes=2, dried=True, dry_program='Sehr Trocken', submission_id='1', used_at=now)
    record_use(resident.member, washes=1, dried=False, submission_id='2', used_at=now)
    voided, _ = record_use(user.member, washes=1, dried=False, submission_id='3', used_at=now)
    voided.voided_at = now
    db.session.commit()
    local = to_local(now)
    rows, totals = monthly_summary(local.year, local.month)
    assert totals == {'entries': 2, 'washes': 3, 'dries': 1, 'total_cents': 75 + 150}
    kassenwart = make_user(email='kw@example.org', name='Kai Kasse', perms=['laundry.view_all'])
    login(kassenwart)
    month = f'{local:%Y-%m}'
    page = client.get(f'/waesche/abrechnung?monat={month}').get_data(as_text=True)
    assert 'Rita Resident' in page and '2,25 €' in page
    response = client.get(f'/waesche/abrechnung.xlsx?monat={month}')
    assert response.status_code == 200
    import io

    wb = load_workbook(io.BytesIO(response.data))
    assert wb.sheetnames == ['Abrechnung', 'Liste']
    assert wb['Abrechnung'].max_row == 4  # header, 2 people, sum
    assert wb['Liste'].max_row == 4  # header and 3 entries (one marked storniert)


def test_manager_voids_with_reason(client, login):
    user = make_resident()
    use, _ = record_use(user.member, washes=1, dried=False, submission_id='v')
    db.session.commit()
    login(make_user(email='con@example.org', name='Con Senior', perms=['laundry.view_all', 'laundry.machines.manage']))
    client.post(f'/waesche/{use.id}/stornieren', data={'reason': 'doppelt'})
    db.session.refresh(use)
    assert use.is_void and use.void_reason == 'doppelt'


def test_settings_price_and_defect_machine(client, login):
    manager = make_resident(email='con@example.org', name='Con Senior')
    from .factories import give_permissions

    give_permissions(manager.member, 'laundry.machines.manage')
    login(manager)
    client.post(
        '/waesche/verwaltung',
        data={
            'price-tier': 'member',
            'price-kind': 'dry',
            'price-price': '0.30',
            'price-valid_from': '2026-11-01',
            'price-submit': 'y',
        },
    )
    price = db.session.scalar(select(LaundryPrice).where(LaundryPrice.valid_from == date(2026, 11, 1)))
    assert (price.tier, price.kind, price.price_cents) == ('member', 'dry', 30)
    dryer = db.session.scalar(select(LaundryMachine).where(LaundryMachine.kind == 'dry'))
    client.post(f'/waesche/verwaltung/maschinen/{dryer.id}', data={'status': 'defect'})
    assert 'Trockner ist als defekt gemeldet' in client.get('/waesche/neu').get_data(as_text=True)


def test_qr_code_leads_to_the_form_and_can_be_rotated(client, login):
    qr = tokens_for('laundry')[0]
    old = qr.token
    response = client.get(f'/q/{old}')
    assert response.status_code == 302 and response.location.endswith('/waesche/neu')
    assert client.get('/q/unknown').status_code == 404
    manager = make_user(email='con@example.org', name='Con Senior', perms=['laundry.machines.manage'])
    login(manager)
    assert 'Wäsche eintragen' in client.get(f'/waesche/verwaltung/qr/{qr.id}').get_data(as_text=True)
    client.post(f'/waesche/verwaltung/qr/{qr.id}/neu')
    db.session.refresh(qr)
    assert qr.token != old
    assert client.get(f'/q/{old}').status_code == 404


def test_month_bounds_follow_berlin_time():
    start, end = month_bounds(2026, 10)
    assert start == datetime(2026, 9, 30, 22, 0, tzinfo=UTC)  # 1 October 00:00 summer time
    assert end == datetime(2026, 10, 31, 23, 0, tzinfo=UTC)  # 1 November 00:00 winter time


def _google_form_file(path, rows):
    wb = Workbook()
    wb.active.title = 'FINANZ'
    ws = wb.create_sheet('LISTE')
    ws.append(['Timestamp', 'Waschen', 'Trocknen', 'Trocknen Modus', 'Transport', 'Email address', 'Timer ', 'Zimmer'])
    for row in rows:
        ws.append(row)
    wb.save(path)
    return path


def test_import_of_the_old_google_form_list(app, tmp_path):
    make_resident(email='known@example.org', name='Known Person')
    path = _google_form_file(
        tmp_path / 'liste.xlsx',
        [
            [
                datetime(2026, 2, 1, 19, 29, 44),
                'beide',
                'Ja!',
                'Schranktrocken',
                'hellblaue wanne',
                'Known@example.org',
                'X',
                None,
            ],
            [datetime(2026, 2, 2, 11, 0, 0), 'eine', None, None, None, None, None, None],
            [datetime(2026, 2, 3, 12, 0, 0), 'eine', None, None, None, 'stranger@example.org', None, None],
        ],
    )
    preview = import_google_form_xlsx(path, dry_run=True)
    assert preview['imported'] == 1 and entries() == []
    stats = import_google_form_xlsx(path, dry_run=False)
    assert (stats['imported'], stats['no_email'], stats['unknown_email']) == (1, 1, {'stranger@example.org': 1})
    (use,) = entries()
    assert (use.washes, use.dried, use.dry_program, use.container, use.source) == (
        2,
        True,
        'Schranktrocken',
        'hellblaue wanne',
        'import',
    )
    assert import_google_form_xlsx(path, dry_run=False)['duplicates'] == 1
    assert len(entries()) == 1
